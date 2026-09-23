#!/usr/bin/env python3
"""
NetJet cold outreach sender for Sduduzo Cele / Charisma Technology.

Reads leads from the "Email Leads" sheet of NetJet_Prospecting_Master.xlsx,
personalises and sends cold outreach emails through mail.charismatech.co.za,
with randomised delays between sends, a POPIA opt out line on every email,
and a CSV send log that later powers threaded replies.

Run this on your own computer, not in a hosted chat environment. It needs
openpyxl installed:

    pip install openpyxl

USAGE

  Render every email locally with no network call and no password, so you
  can read them before anything goes out:

    python netjet_outreach.py preview --input NetJet_Prospecting_Master.xlsx

  Send a test batch. Every email in this mode goes to your own inbox
  instead of the real recipient, so you can see exactly what a prospect
  would receive:

    python netjet_outreach.py send --input NetJet_Prospecting_Master.xlsx --test-limit 5

  Send for real, once you are happy with the test:

    python netjet_outreach.py send --input NetJet_Prospecting_Master.xlsx --live --limit 50

  Running the same live command again later picks up where the last run
  stopped: anyone already emailed live (per the send log) is skipped.

  Someone replied STOP? Add them to the suppression list so they are never
  emailed again:

    python netjet_outreach.py stop someone@company.co.za

  Reply to someone who wrote back, threaded onto your original email:

    python netjet_outreach.py reply --to someone@company.co.za --body-file reply.txt

Your mailbox password is asked for at the start of any run that needs it.
It is typed hidden, held in memory only for that run, and never written to
this file, the log, or anywhere on disk.
"""

import argparse
import csv
import getpass
import os
import random
import re
import smtplib
import ssl
import sys
import time
from datetime import datetime, timezone
from email.mime.text import MIMEText
from email.utils import make_msgid, formatdate

import openpyxl

# ---------------------------------------------------------------------------
# Configuration. Edit these if any detail changes.
# ---------------------------------------------------------------------------

SMTP_HOST = "mail.charismatech.co.za"
SMTP_PORT = 465

IMAP_HOST = "mail.charismatech.co.za"
IMAP_PORT = 993
# cPanel mailboxes usually call the Sent folder one of these. If the reply
# command can't find a Sent match, try changing this.
IMAP_SENT_FOLDER = '"Sent"'

FROM_EMAIL = "sduduzo@charismatech.co.za"
FROM_NAME = "Sduduzo Cele"

SIGNATURE = (
    "Kind regards,\n"
    "Sduduzo Cele\n"
    "Sales Representative, Charisma Technology\n"
    "+27 71 709 8160\n"
    "WhatsApp +27 67 872 9242\n"
    "sduduzo@charismatech.co.za"
)

OPT_OUT_LINE = (
    "If you would rather not receive outreach emails from us, reply with "
    "the word STOP and we will remove your details from this list."
)

MIN_DELAY_SECONDS = 20
MAX_DELAY_SECONDS = 45

SEND_LOG_PATH = "netjet_send_log.csv"
LOG_FIELDNAMES = [
    "timestamp", "company", "email", "contact_name", "subject",
    "template_index", "message_id", "status", "error", "mode",
]

# One email address per line. Anyone listed here is never emailed. Add to it
# with the "stop" command whenever someone replies STOP.
SUPPRESSION_PATH = "netjet_suppress.txt"

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")

HONORIFICS = {"mr", "mrs", "ms", "miss", "dr", "prof", "mnr", "mev"}

# ---------------------------------------------------------------------------
# Email copy. Three full variants so the 50 don't read identically. Each
# lead gets one, rotated in order.
# ---------------------------------------------------------------------------

TEMPLATES = [
    {
        "subject": "Stop losing margin on projects you already won",
        "body": """Hi {greeting_name},

Quick question. When a project runs over budget, do you find out before the invoice does, or after?

Most project based businesses only find out once the timesheets are captured, approved and finally reconciled weeks later. By then the margin is already gone.

I work with NetJet, a cloud based team and time management platform built by Charisma Technology, a company with 15 years of consulting and custom software experience in Durban. They built NetJet because they needed it themselves, and they still run their own consulting business on it today.

What it gives you:
Live time tracking and timesheet approval, so you see hours against budget as they happen
Project admin, quoting and purchase orders in one place
Billing and dashboards that show project profitability at a glance
Sage integration if you already run your accounts there

Pricing is simple. R149 per user a month for the full toolset, or R99 per user a month if you only need time tracking to start.

Would you be open to a short call this week? I can show you exactly how it would sit on top of your current projects, no obligation.

{signature}

{opt_out}""",
    },
    {
        "subject": "Missed another deadline this month?",
        "body": """Hi {greeting_name},

If a project on your books is running late right now, you are not alone. Most construction and project businesses lose time not on the work itself, but on the admin around it, chasing timesheets, redoing quotes, and finding out about scope creep after the fact.

I represent NetJet, a cloud based project and time management platform out of Durban, built and used daily by Charisma Technology in their own consulting business over the last 15 years.

With NetJet your team logs time and progress as they work, your project leads see status and budget in real time, and quoting, purchase orders and billing all run from the same system instead of five different spreadsheets.

We offer it at R149 per user a month for full access, or R99 per user a month for a lighter time tracking only option if you want to start small.

I would like to show you a short demo built around a project like yours. Would you have 15 minutes this week or next?

{signature}

{opt_out}""",
    },
    {
        "subject": "A faster way to see project costs at {company}",
        "body": """Hi {greeting_name},

Most of the project businesses I speak to run timesheets, quoting and billing across two or three disconnected tools, and only see the real cost of a job once it is finished.

NetJet brings that into one system. Your team logs hours and progress against each project, your office sees budget and profitability while the job is still running, and quoting, purchase orders and billing all sit in the same place. It is built and used every day by Charisma Technology, a Durban based software company with 15 years in custom project systems.

Two ways to get started. R99 per user a month for time tracking on its own, or R149 per user a month for the full set of tools including quoting, billing and Sage integration.

Happy to run a short demo using a project similar to what {company} handles, so you can see it against real numbers rather than a generic pitch. Would that be worth 15 minutes?

{signature}

{opt_out}""",
    },
]


# ---------------------------------------------------------------------------
# Loading leads
# ---------------------------------------------------------------------------

def pick_email(raw, name):
    """First address in the cell, or the one whose local part matches the
    contact's first name, so "Hi Hardus" goes to hardus@ rather than info@."""
    if not raw:
        return None
    matches = EMAIL_RE.findall(str(raw))
    if not matches:
        return None
    if name:
        for m in matches:
            if name.lower() in m.split("@")[0].lower():
                return m
    return matches[0]


def greeting_name(raw):
    """Name to put after "Hi", or None to fall back to "<company> team".

    "Adie van Oosten" -> "Adie", "Mr V M Matsheke" -> "Mr Matsheke",
    "JWB Stoltz" -> "JWB Stoltz", "(name not published ...)" -> None.
    """
    if not raw:
        return None
    text = str(raw).strip()
    if not text or text.lower() == "nan" or text.startswith("("):
        return None
    text = re.sub(r"\(.*?\)", "", text).strip()
    parts = text.split()
    if not parts:
        return None
    if parts[0].lower().rstrip(".") in HONORIFICS:
        return f"{parts[0]} {parts[-1]}" if len(parts) > 1 else None
    if parts[0].isupper() and len(parts[0]) <= 3 and len(parts) > 1:
        return " ".join(parts)
    return parts[0]


def load_suppressed():
    if not os.path.exists(SUPPRESSION_PATH):
        return set()
    with open(SUPPRESSION_PATH, encoding="utf-8") as f:
        return {line.strip().lower() for line in f if line.strip()}


def load_already_sent():
    """Addresses that already received a live email, so reruns don't repeat."""
    if not os.path.exists(SEND_LOG_PATH):
        return set()
    with open(SEND_LOG_PATH, encoding="utf-8") as f:
        return {
            r["email"].lower() for r in csv.DictReader(f)
            if r["status"] == "sent" and r["mode"] == "live"
        }


def load_leads(path, sheet_name, include_previously_emailed):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        sys.exit(f"Sheet '{sheet_name}' not found. Sheets in file: {wb.sheetnames}")
    ws = wb[sheet_name]

    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers) if h}

    required = ["Company", "Email", "Previously Emailed (Job Search)"]
    missing = [c for c in required if c not in idx]
    if missing:
        sys.exit(f"Expected column(s) not found in '{sheet_name}': {missing}")

    suppressed = load_suppressed()
    already_sent = load_already_sent()
    seen = set()

    leads = []
    skipped = {
        "previously emailed for the job search": 0,
        "with no usable email address": 0,
        "duplicate email address": 0,
        "on the STOP list": 0,
        "already emailed live": 0,
    }

    for row in ws.iter_rows(min_row=2, values_only=True):
        company = row[idx["Company"]]
        if not company:
            continue

        contact_name = row[idx["Contact Name"]] if "Contact Name" in idx else None
        name = greeting_name(contact_name)

        email = pick_email(row[idx["Email"]], name)
        if not email:
            skipped["with no usable email address"] += 1
            continue

        key = email.lower()
        if key in seen:
            skipped["duplicate email address"] += 1
            continue
        seen.add(key)

        if key in suppressed:
            skipped["on the STOP list"] += 1
            continue

        if key in already_sent:
            skipped["already emailed live"] += 1
            continue

        prev_raw = row[idx["Previously Emailed (Job Search)"]]
        prev_yes = str(prev_raw).strip().lower() == "yes"
        if prev_yes and not include_previously_emailed:
            skipped["previously emailed for the job search"] += 1
            continue

        leads.append({
            "company": " ".join(str(company).split()),
            "email": email,
            "contact_name": name,
            "previously_emailed": prev_yes,
        })

    return leads, skipped


def describe_skipped(skipped):
    parts = [f"{n} {why}" for why, n in skipped.items() if n]
    return "Skipped " + ", ".join(parts) + "." if parts else "Nothing skipped."


# ---------------------------------------------------------------------------
# Building messages
# ---------------------------------------------------------------------------

def build_message(lead, position):
    template_index = position % len(TEMPLATES)
    template = TEMPLATES[template_index]
    greeting = lead["contact_name"] or f"{lead['company']} team"
    subject = template["subject"].format(company=lead["company"])
    body = template["body"].format(
        greeting_name=greeting,
        company=lead["company"],
        signature=SIGNATURE,
        opt_out=OPT_OUT_LINE,
    )
    return subject, body, template_index


# ---------------------------------------------------------------------------
# Send log
# ---------------------------------------------------------------------------

def append_log(rows):
    if not rows:
        return
    new_file = not os.path.exists(SEND_LOG_PATH)
    with open(SEND_LOG_PATH, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=LOG_FIELDNAMES)
        if new_file:
            writer.writeheader()
        writer.writerows(rows)


def find_in_log(to_email):
    if not os.path.exists(SEND_LOG_PATH):
        return None
    with open(SEND_LOG_PATH, encoding="utf-8") as f:
        matches = [
            r for r in csv.DictReader(f)
            if r["email"].lower() == to_email.lower()
            and r["status"] == "sent" and r["mode"] == "live"
        ]
    return matches[-1] if matches else None


# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------

def smtp_connect(password):
    smtp = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ssl.create_default_context())
    smtp.login(FROM_EMAIL, password)
    return smtp


def cmd_preview(args):
    leads, skipped = load_leads(args.input, args.sheet, args.include_previously_emailed)
    total = len(leads)
    if args.limit:
        leads = leads[: args.limit]

    shown = f" (showing first {len(leads)})" if args.limit and len(leads) < total else ""
    print(f"{total} leads would receive an email{shown}. {describe_skipped(skipped)}\n")

    for i, lead in enumerate(leads):
        subject, body, template_index = build_message(lead, i)
        print("=" * 72)
        print(f"To: {lead['company']} <{lead['email']}>  (template {template_index + 1})")
        print(f"Subject: {subject}")
        print("-" * 72)
        print(body)
        print()


def cmd_send(args):
    leads, skipped = load_leads(args.input, args.sheet, args.include_previously_emailed)

    if not args.live and args.test_limit:
        leads = leads[: args.test_limit]
    elif args.limit:
        leads = leads[: args.limit]

    print(f"{len(leads)} leads loaded. {describe_skipped(skipped)}")
    if not leads:
        print("Nothing to send.")
        return

    if not args.live:
        print(
            f"TEST MODE: all {len(leads)} emails will be sent to {FROM_EMAIL} "
            "(your own inbox), not to the real recipients. Pass --live to send for real."
        )
    else:
        print(f"LIVE MODE: {len(leads)} real emails are about to go out. Ctrl+C now to stop.")

    password = getpass.getpass(f"Password for {FROM_EMAIL}: ")
    smtp = smtp_connect(password)
    sent_count = 0

    try:
        for i, lead in enumerate(leads):
            subject, body, template_index = build_message(lead, i)

            msg = MIMEText(body, "plain", "utf-8")
            msg["Subject"] = subject
            msg["From"] = f"{FROM_NAME} <{FROM_EMAIL}>"
            msg["To"] = FROM_EMAIL if not args.live else lead["email"]
            msg["Date"] = formatdate(localtime=True)
            msg_id = make_msgid(domain="charismatech.co.za")
            msg["Message-ID"] = msg_id

            status, error = "sent", ""
            try:
                try:
                    smtp.send_message(msg)
                except smtplib.SMTPServerDisconnected:
                    # Long runs can outlive the server's session. Reconnect once.
                    smtp = smtp_connect(password)
                    smtp.send_message(msg)
                sent_count += 1
                tag = "LIVE" if args.live else "TEST -> self"
                print(f"[{i + 1}/{len(leads)}] Sent to {lead['company']} <{lead['email']}> ({tag})")
            except Exception as exc:
                status, error = "failed", str(exc)
                print(f"[{i + 1}/{len(leads)}] FAILED for {lead['company']} <{lead['email']}>: {exc}")

            # Written after every send, so a crash or Ctrl+C never loses a
            # record of who was already emailed.
            append_log([{
                "timestamp": datetime.now(timezone.utc).isoformat(),
                "company": lead["company"],
                "email": lead["email"],
                "contact_name": lead["contact_name"] or "",
                "subject": subject,
                "template_index": template_index,
                "message_id": msg_id,
                "status": status,
                "error": error,
                "mode": "live" if args.live else "test",
            }])

            if i < len(leads) - 1:
                time.sleep(random.uniform(MIN_DELAY_SECONDS, MAX_DELAY_SECONDS))
    except KeyboardInterrupt:
        print("\nStopped. Everything sent so far is in the log; rerun to continue.")
    finally:
        try:
            smtp.quit()
        except Exception:
            pass

    print(f"Done. {sent_count} sent, log at {SEND_LOG_PATH}")


def cmd_stop(args):
    suppressed = load_suppressed()
    added = [e for e in args.emails if e.strip().lower() not in suppressed]
    with open(SUPPRESSION_PATH, "a", encoding="utf-8") as f:
        for e in added:
            f.write(e.strip().lower() + "\n")
    print(f"Added {len(added)} address(es) to {SUPPRESSION_PATH}. They will not be emailed again.")


def cmd_reply(args):
    body_text = args.body
    if args.body_file:
        with open(args.body_file, encoding="utf-8") as f:
            body_text = f.read()
    if not body_text:
        sys.exit("Provide reply text with --body or --body-file.")

    password = None
    original = find_in_log(args.to)

    if original is None:
        print(f"No earlier live email to {args.to} found in {SEND_LOG_PATH}. Checking Sent folder over IMAP...")
        password = getpass.getpass(f"Password for {FROM_EMAIL}: ")
        original = find_via_imap(args.to, password)

    if original is None:
        print(f"Could not find an earlier outreach email to {args.to}. Nothing sent.")
        return

    subject = original["subject"]
    if not subject.lower().startswith("re:"):
        subject = "Re: " + subject

    if password is None:
        password = getpass.getpass(f"Password for {FROM_EMAIL}: ")

    msg = MIMEText(f"{body_text}\n\n{SIGNATURE}", "plain", "utf-8")
    msg["Subject"] = subject
    msg["From"] = f"{FROM_NAME} <{FROM_EMAIL}>"
    msg["To"] = args.to
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid(domain="charismatech.co.za")
    msg["In-Reply-To"] = original["message_id"]
    msg["References"] = original["message_id"]

    smtp = smtp_connect(password)
    try:
        smtp.send_message(msg)
    finally:
        smtp.quit()

    print(f"Reply sent to {args.to}, threaded onto {original['message_id']}.")


def find_via_imap(to_email, password):
    import imaplib
    import email as email_lib

    if not EMAIL_RE.fullmatch(to_email):
        print(f"'{to_email}' does not look like an email address.")
        return None

    imap = None
    try:
        imap = imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT)
        imap.login(FROM_EMAIL, password)
        imap.select(IMAP_SENT_FOLDER, readonly=True)
        typ, data = imap.search(None, f'(TO "{to_email}")')
        if typ != "OK" or not data or not data[0]:
            return None
        latest_id = data[0].split()[-1]
        typ, msg_data = imap.fetch(latest_id, "(RFC822)")
        parsed = email_lib.message_from_bytes(msg_data[0][1])
        return {"subject": parsed.get("Subject", ""), "message_id": parsed.get("Message-ID", "")}
    except Exception as exc:
        print(f"IMAP lookup failed: {exc}")
        return None
    finally:
        if imap is not None:
            try:
                imap.logout()
            except Exception:
                pass


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description="NetJet cold outreach sender")
    sub = parser.add_subparsers(dest="command", required=True)

    p_preview = sub.add_parser("preview", help="Render every email locally. No network, no password.")
    p_preview.add_argument("--input", default="NetJet_Prospecting_Master.xlsx")
    p_preview.add_argument("--sheet", default="Email Leads")
    p_preview.add_argument("--limit", type=int, default=None)
    p_preview.add_argument("--include-previously-emailed", action="store_true")
    p_preview.set_defaults(func=cmd_preview)

    p_send = sub.add_parser("send", help="Send emails over SMTP.")
    p_send.add_argument("--input", default="NetJet_Prospecting_Master.xlsx")
    p_send.add_argument("--sheet", default="Email Leads")
    p_send.add_argument("--limit", type=int, default=None, help="Cap on live sends.")
    p_send.add_argument("--test-limit", type=int, default=5, help="How many test emails to send to yourself when --live is not set.")
    p_send.add_argument("--include-previously-emailed", action="store_true")
    p_send.add_argument("--live", action="store_true", help="Send to real recipients. Without this, everything goes to your own inbox.")
    p_send.set_defaults(func=cmd_send)

    p_stop = sub.add_parser("stop", help="Add addresses that replied STOP to the suppression list.")
    p_stop.add_argument("emails", nargs="+")
    p_stop.set_defaults(func=cmd_stop)

    p_reply = sub.add_parser("reply", help="Send a threaded reply to a lead who wrote back.")
    p_reply.add_argument("--to", required=True)
    p_reply.add_argument("--body")
    p_reply.add_argument("--body-file")
    p_reply.set_defaults(func=cmd_reply)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
